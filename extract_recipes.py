import openpyxl
import re
import json

wb = openpyxl.load_workbook('RM 2019_v2.xlsx', data_only=True)

SHEET_CONFIG = {
    'Beverages': {'name_col': 2, 'qty_col': 3, 'unit_col': 6, 'notes_col': 7},
    'Breakfast': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6, 'stage_col': 1},
    'Chutneys': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6, 'stage_col': 1},
    'Dal': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6},
    'Economical Recipes': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Ekadashi': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Fried Items': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6},
    'Jam': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6},
    'Kadhi': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Khichadi': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Masalas': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Rice': {'name_col': 1, 'qty_col': 2, 'unit_col': 5, 'notes_col': 6},
    'Roti': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Soups': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Sweets': {'name_col': 2, 'qty_col': 3, 'unit_col': 6, 'notes_col': 7},
    'Sweets Sustainable': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Subji, Dry': {'name_col': 1, 'qty_col': 2, 'unit_col': 4, 'notes_col': 5},
    'Wet Subjis': {'name_col': 2, 'qty_col': 3, 'unit_col': 6, 'notes_col': 7},
    'FHC Sabjis': {'name_col': 2, 'qty_col': 3, 'unit_col': 5, 'notes_col': 6},
}

def clean_str(val):
    if val is None:
        return ''
    return str(val).strip()

def clean_num(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return round(float(val), 4)
    s = str(val).strip().replace(',', '')
    try:
        return round(float(s), 4)
    except:
        return None

def is_yield_cell(val):
    if not val:
        return False
    s = str(val).strip().lower()
    # It must start with yield or recipe yield, or be batter yield, or yield by vol/wt
    if s.startswith('yield') or s.startswith('recipe yield') or s.startswith('batter yield'):
        return True
    return False

def parse_sheet_recipes_refined(sheet_name):
    if sheet_name not in SHEET_CONFIG:
        return []
    cfg = SHEET_CONFIG[sheet_name]
    name_col = cfg['name_col']
    qty_col = cfg['qty_col']
    unit_col = cfg.get('unit_col', 5)
    notes_col = cfg.get('notes_col', 6)
    stage_col = cfg.get('stage_col', None)
    
    ws = wb[sheet_name]
    
    # Locate yield rows
    yield_groups = [] # list of lists of consecutive yield row indices
    r = 1
    while r <= ws.max_row:
        v = ws.cell(r, name_col).value
        if not v and name_col > 1:
            v = ws.cell(r, name_col - 1).value
            
        if is_yield_cell(v):
            group = [r]
            # check if following rows are also yield rows
            r_next = r + 1
            while r_next <= ws.max_row:
                v_next = ws.cell(r_next, name_col).value
                if not v_next and name_col > 1:
                    v_next = ws.cell(r_next, name_col - 1).value
                if is_yield_cell(v_next):
                    group.append(r_next)
                    r_next += 1
                else:
                    break
            yield_groups.append(group)
            r = r_next
        else:
            r += 1

    recipes = []
    for g_idx, y_group in enumerate(yield_groups):
        first_yr = y_group[0]
        last_yr = y_group[-1]
        
        # Look above first_yr for title
        title = ''
        cuisine = ''
        for r_search in range(first_yr - 1, max(0, first_yr - 5), -1):
            val = clean_str(ws.cell(r_search, name_col).value)
            if not val and name_col > 1:
                val = clean_str(ws.cell(r_search, 1).value)
            if val and val.lower() not in ['back', 'home', ''] and not is_yield_cell(val):
                title = val
                for extra_c in range(name_col + 1, min(name_col + 8, ws.max_column + 1)):
                    c_val = clean_str(ws.cell(r_search, extra_c).value)
                    if c_val and c_val.lower() not in ['back', 'home', 'rate', 'amount']:
                        cuisine = c_val
                        break
                break
                
        if not title:
            # check the yield label itself e.g. "Yield (Big size Vada)"
            y_label = clean_str(ws.cell(first_yr, name_col).value)
            title = f"{sheet_name} Recipe {g_idx+1}"
            
        # Parse yield from first yield row
        base_yield = clean_num(ws.cell(first_yr, qty_col).value) or 100.0
        base_unit = ''
        for u_c in [unit_col, unit_col - 1, unit_col + 1, qty_col + 1, qty_col + 2]:
            u_val = clean_str(ws.cell(first_yr, u_c).value)
            if u_val and u_val.lower() not in ['yield', 'target', 'base', 'rate', 'amount'] and len(u_val) < 10:
                base_unit = u_val
                break
        if not base_unit:
            base_unit = 'L' if any(w in sheet_name.lower() for w in ['soup', 'dal', 'bever', 'kadhi']) else 'Kg'
            
        # Next group starts at
        next_yr_limit = yield_groups[g_idx + 1][0] if g_idx + 1 < len(yield_groups) else ws.max_row + 1
        
        ingredients = []
        instructions = []
        current_stage = 'Main'
        
        for row_idx in range(last_yr + 1, next_yr_limit):
            item_name = clean_str(ws.cell(row_idx, name_col).value)
            item_qty = clean_num(ws.cell(row_idx, qty_col).value)
            
            if stage_col:
                s_val = clean_str(ws.cell(row_idx, stage_col).value)
                if s_val and s_val.lower() not in ['back', 'home']:
                    current_stage = s_val
                    
            if item_name.lower() in ['back', 'home', 'compiled on', 'last edited on']:
                # Read following rows for instructions
                for r_after in range(row_idx, min(row_idx + 12, next_yr_limit)):
                    instr_val = clean_str(ws.cell(r_after, name_col).value)
                    if instr_val and instr_val.lower() not in ['back', 'home', 'compiled on', 'last edited on']:
                        instructions.append(instr_val)
                break
                
            if not item_name:
                row_text = ' '.join([clean_str(ws.cell(row_idx, c).value) for c in range(1, 10)])
                if len(row_text) > 20 and not any(k in row_text.lower() for k in ['back', 'home', 'compiled on']):
                    instructions.append(row_text)
                continue
                
            item_unit = ''
            for u_c in [unit_col, unit_col - 1, unit_col + 1, qty_col + 1, qty_col + 2]:
                u_val = clean_str(ws.cell(row_idx, u_c).value)
                if u_val and len(u_val) <= 8 and not any(char.isdigit() for char in u_val):
                    item_unit = u_val
                    break
                    
            item_notes = ''
            for n_c in [notes_col, notes_col + 1, notes_col + 2]:
                n_val = clean_str(ws.cell(row_idx, n_c).value)
                if n_val and len(n_val) > 1 and n_val != item_unit:
                    item_notes = n_val
                    break
                    
            if item_qty is not None or (item_name and not item_name.startswith('---')):
                ingredients.append({
                    'name': item_name,
                    'stage': current_stage,
                    'quantity': item_qty if item_qty is not None else 0.0,
                    'unit': item_unit or 'gm',
                    'notes': item_notes
                })
                
        # Only accept if it has ingredients!
        if len(ingredients) > 0:
            # Clean title
            if title.strip().startswith('1.') or title.strip().startswith('2.'):
                # note that got mistaken for title
                pass
            recipes.append({
                'id': f"{sheet_name.lower().replace(' ', '_').replace(',', '')}_{len(recipes)+1}",
                'title': title,
                'category': sheet_name,
                'cuisine': cuisine,
                'base_yield': base_yield,
                'yield_unit': base_unit,
                'ingredients_count': len(ingredients),
                'ingredients': ingredients,
                'instructions': instructions
            })
            
    return recipes

all_refined = []
for sname in SHEET_CONFIG.keys():
    recs = parse_sheet_recipes_refined(sname)
    all_refined.extend(recs)
    print(f"Sheet {sname:20s}: {len(recs)} valid recipes")

print(f"=== TOTAL CLEAN RECIPES EXTRACTED: {len(all_refined)} ===")

with open('recipes_db.json', 'w', encoding='utf-8') as f:
    json.dump(all_refined, f, indent=2, ensure_ascii=False)
print("Updated recipes_db.json!")
